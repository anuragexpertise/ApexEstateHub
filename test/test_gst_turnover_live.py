"""
Job 3 -- GST aggregate turnover (CGST Act s.2(6)): a classified worksheet and a three-state gate.

Live tests against a REAL Postgres, rolled back.

    ESTATEHUB_LIVE_TESTS=1 PGHOST=... PGDATABASE=... PGUSER=... pytest test/test_gst_turnover_live.py
"""
import os
from datetime import date

import pytest

from test.live_db_gate import LIVE_DB_REASON, live_db_enabled

psycopg2 = pytest.importorskip("psycopg2")
pytestmark = pytest.mark.skipif(not live_db_enabled(), reason=LIVE_DB_REASON)

SOC = 1


def _fy(d=None):
    d = d or date.today()
    return d.year if d.month >= 4 else d.year - 1


FY = _fy()


@pytest.fixture()
def conn():
    c = psycopg2.connect(
        host=os.getenv("PGHOST"), port=os.getenv("PGPORT", "5432"),
        dbname=os.getenv("PGDATABASE"), user=os.getenv("PGUSER"),
        password=os.getenv("PGPASSWORD"), sslmode=os.getenv("PGSSLMODE", "prefer"),
    )
    cur = c.cursor()
    cur.execute("SELECT to_regproc('fn_gst_turnover_gate') IS NOT NULL")
    if not cur.fetchone()[0]:
        c.close()
        pytest.skip("estatehub.sql with the s.2(6) turnover worksheet is not installed in this database")
    cur.execute("SET CONSTRAINTS ALL IMMEDIATE")
    yield c
    c.rollback()
    c.close()


@pytest.fixture()
def cur(conn):
    return conn.cursor()


def one(cur, sql, args=()):
    cur.execute(sql, args)
    return cur.fetchone()[0]


def acc(cur, tab):
    return one(cur, "SELECT id FROM accounts WHERE society_id=%s AND tab_name=%s ORDER BY id LIMIT 1", (SOC, tab))


def post_credit(cur, tab, amount, on=None):
    """Dr Sundry Debtors / Cr the account of this ledger head, in one balanced journal."""
    sdr = one(cur, "SELECT id FROM accounts WHERE society_id=%s AND name ILIKE 'Sundry Debtors' LIMIT 1", (SOC,))
    j = one(cur, "SELECT NEXTVAL('seq_transaction_number')")
    on = on or date.today()
    cur.execute("""INSERT INTO transactions (society_id, entry_side, trx_date, acc_id, amount, mode, status, journal_id)
                   VALUES (%s,'Dr',%s,%s,%s,'journal','paid',%s), (%s,'Cr',%s,%s,%s,'journal','paid',%s)""",
                (SOC, on, sdr, amount, j, SOC, on, acc(cur, tab), amount, j))


def gate(cur):
    cur.execute("SELECT lower_bound, upper_bound, unclassified, threshold, state FROM fn_gst_turnover_gate(%s,%s)", (SOC, FY))
    lo, up, un, thr, st = cur.fetchone()
    return float(lo), float(up), float(un), float(thr), st


def set_threshold(cur, rupees):
    # The table holds the threshold in lakh to 2 decimals (Rs 1,000), so tests leave a wide margin.
    cur.execute("UPDATE state_compliance_thresholds SET value=%s WHERE threshold_key='gst_turnover_lakh'", (round(rupees / 100000.0, 2),))


def test_tax_collected_is_not_turnover(cur):
    lo0, up0, *_ = gate(cur)
    post_credit(cur, "CGST", 500000)
    lo1, up1, *_ = gate(cur)
    assert (lo1, up1) == (lo0, up0)          # s.2(6) excludes central/state tax


def test_balance_sheet_credits_are_not_turnover(cur):
    lo0, up0, *_ = gate(cur)
    post_credit(cur, "CapAc", 800000)         # capital account
    post_credit(cur, "Eq", 300000)            # reserves
    lo1, up1, *_ = gate(cur)
    assert (lo1, up1) == (lo0, up0)


def test_exempt_supplies_are_included(cur):
    lo0, up0, *_ = gate(cur)
    post_credit(cur, "IncExmpt", 250000)
    lo1, up1, *_ = gate(cur)
    assert lo1 - lo0 == 250000 and up1 - up0 == 250000     # the old function dropped these


def test_member_charges_are_taxable_supplies(cur):
    lo0, *_ = gate(cur)
    post_credit(cur, "SocM", 400000)
    assert gate(cur)[0] - lo0 == 400000


def test_a_contested_head_is_unclassified_not_guessed(cur):
    lo0, up0, un0, *_ = gate(cur)
    post_credit(cur, "IntFD", 600000)         # interest on deposits
    lo1, up1, un1, *_ = gate(cur)
    assert lo1 == lo0 and up1 - up0 == 600000 and un1 - un0 == 600000


def test_three_states_and_classification_resolves_doubt(cur):
    post_credit(cur, "SocM", 100000)
    post_credit(cur, "IntFD", 300000)
    lo, up, un, thr, st = gate(cur)
    set_threshold(cur, lo + 100000)           # between lower and upper
    assert gate(cur)[4] == "indeterminate"
    set_threshold(cur, lo - 50000)
    assert gate(cur)[4] == "above"            # classified supplies alone exceed it
    set_threshold(cur, up + 50000)
    assert gate(cur)[4] == "below"            # even counting everything it is within it
    # the CA classifies the interest as exempt supply: it now counts, and the doubt is gone
    set_threshold(cur, lo + 100000)
    cur.execute("SELECT fn_set_gst_supply_class(%s,%s,'exempt',%s)", (SOC, acc(cur, "IntFD"), "CA advice dated 2026-10-01"))
    assert gate(cur)[4] == "above" and gate(cur)[0] == lo + 300000
    # or as not-a-supply: it drops out, and the society is below the threshold
    cur.execute("SELECT fn_set_gst_supply_class(%s,%s,'not_a_supply',%s)", (SOC, acc(cur, "IntFD"), "CA advice dated 2026-10-01"))
    assert gate(cur)[4] == "below"


def test_a_non_taxable_class_needs_a_reference(cur):
    cur.execute("SAVEPOINT s")
    with pytest.raises(psycopg2.Error):
        cur.execute("SELECT fn_set_gst_supply_class(%s,%s,'exempt','')", (SOC, acc(cur, "IntFD")))
    cur.execute("ROLLBACK TO SAVEPOINT s")


def test_legacy_exempt_flag_still_counts_as_exempt(cur):
    cur.execute("UPDATE accounts SET is_gst_exempt_supply = TRUE WHERE society_id=%s AND tab_name='IntBK'", (SOC,))
    post_credit(cur, "IntBK", 50000)
    cur.execute("SELECT supply_class, class_source FROM fn_gst_turnover_lines(%s,%s) WHERE tab_name='IntBK'", (SOC, FY))
    assert cur.fetchone() == ("exempt", "legacy flag")


def test_old_single_number_function_is_the_classified_lower_bound(cur):
    post_credit(cur, "SocM", 123456)
    assert float(one(cur, "SELECT fn_society_turnover_fy(%s,%s)", (SOC, FY))) == gate(cur)[0]


def test_worksheet_is_kept_and_a_confirmed_one_is_a_record(cur):
    post_credit(cur, "SocM", 100000)
    wid = one(cur, "SELECT fn_save_gst_turnover_worksheet(%s,%s,NULL)", (SOC, FY))
    cur.execute("SELECT gate_state, lower_bound, jsonb_array_length(lines) FROM gst_turnover_worksheets WHERE id=%s", (wid,))
    state, lower, n = cur.fetchone()
    assert state in ("above", "below", "indeterminate") and float(lower) == gate(cur)[0] and n >= 1
    cur.execute("SAVEPOINT s")
    with pytest.raises(psycopg2.Error):
        cur.execute("SELECT fn_confirm_gst_turnover_worksheet(%s,NULL,'x')", (wid,))      # needs a real reference
    cur.execute("ROLLBACK TO SAVEPOINT s")
    uid = one(cur, "SELECT id FROM users WHERE society_id=%s LIMIT 1", (SOC,))
    cur.execute("SELECT fn_confirm_gst_turnover_worksheet(%s,%s,'CA certificate GST/2026/17')", (wid, uid))
    for stmt in ("UPDATE gst_turnover_worksheets SET threshold=1 WHERE id=%s", "DELETE FROM gst_turnover_worksheets WHERE id=%s"):
        cur.execute("SAVEPOINT s")
        with pytest.raises(psycopg2.Error):
            cur.execute(stmt, (wid,))
        cur.execute("ROLLBACK TO SAVEPOINT s")


# ---- billing: GST is charged only when the gate says 'above' -----------------------------------------------
def _bill(cur):
    cur.execute("SET CONSTRAINTS ALL DEFERRED")      # billing posts each journal's legs in separate statements
    cur.execute("DELETE FROM receivables WHERE society_id=%s", (SOC,))
    cur.execute("DELETE FROM apt_charges_fines_basis WHERE society_id=%s", (SOC,))
    cur.execute("SELECT fn_backfill_undivided_interest(%s, TRUE)", (SOC,))
    cur.execute("""INSERT INTO apt_charges_fines_basis (society_id, apt_id, apt_maintenance_amount, apt_maintenance_rate,
                       apt_due_day, apt_interest_pct, start_date, apt_status, billing_basis, common_expense_budget_monthly)
                   VALUES (%s, NULL, 0, 3.0, 5, 0, DATE_TRUNC('month', CURRENT_DATE)::DATE, TRUE, 'undivided_interest', 5000000)""", (SOC,))
    cur.execute("UPDATE apartments SET apt_calc_start_date = DATE_TRUNC('month', CURRENT_DATE)::DATE WHERE society_id=%s", (SOC,))
    cur.execute("SELECT fn_auto_generate_receivables(%s)", (SOC,))
    return one(cur, "SELECT count(*) FROM receivables WHERE society_id=%s AND description ILIKE 'CGST%%'", (SOC,))


def test_billing_charges_gst_only_when_the_gate_is_above(cur):
    post_credit(cur, "SocM", 100000)
    lo, up, *_ = gate(cur)
    set_threshold(cur, lo - 50000)
    assert gate(cur)[4] == "above"
    assert _bill(cur) > 0


def test_billing_holds_gst_and_raises_a_flag_when_indeterminate(cur):
    post_credit(cur, "SocM", 100000)
    post_credit(cur, "IntFD", 300000)
    lo, up, *_ = gate(cur)
    set_threshold(cur, lo + 100000)
    assert gate(cur)[4] == "indeterminate"
    assert _bill(cur) == 0                   # tax that may not be due is not collected
    assert one(cur, "SELECT count(*) FROM compliance_flags WHERE society_id=%s AND rule_code='GST_TURNOVER_INDETERMINATE'", (SOC,)) == 1


def test_billing_does_not_charge_gst_when_below(cur):
    lo, up, *_ = gate(cur)
    set_threshold(cur, up + 50000)
    assert gate(cur)[4] == "below"
    assert _bill(cur) == 0


def test_export_carries_the_worksheet_and_turnover_excludes_tax(cur):
    import io
    from openpyxl import load_workbook
    from database import gst_export

    class Shim:
        def _execute(self, sql, params=None, fetch_one=False, fetch_all=False):
            cur.execute(sql, params)
            if fetch_one:
                r = cur.fetchone()
                return None if r is None else {d[0]: v for d, v in zip(cur.description, r)}
            if fetch_all:
                return [{d[0]: v for d, v in zip(cur.description, r)} for r in cur.fetchall()]

    post_credit(cur, "SocM", 100000)
    post_credit(cur, "IntFD", 50000)
    wb = load_workbook(io.BytesIO(gst_export.generate_gst_summary_excel(Shim(), SOC, FY)))
    assert "Aggregate Turnover" in wb.sheetnames
    cells = [c.value for row in wb["Aggregate Turnover"].iter_rows() for c in row if c.value is not None]
    assert "UNCLASSIFIED" in cells and any(str(v).startswith("INDETERMINATE") or str(v).startswith(("ABOVE", "BELOW")) for v in cells)
    assert "Value of Supplies (excl. GST)" in [c.value for c in wb["Summary"][2]]
