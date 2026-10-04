"""
UP AOA Compliance card: renderer, callback wiring and write handlers, against a REAL Postgres.

Skipped unless PGHOST is set (same convention as test_up_aoa_compliance_live.py). Handlers
normally go through database.db_manager.db, which commits on its own connections; here db
is swapped for a wrapper on ONE connection that is rolled back at the end of each test, so
the seeded demo data is never changed.
"""
import io
import os
from datetime import date

import pytest

from test.live_db_gate import LIVE_DB_REASON, live_db_enabled

pytestmark = pytest.mark.skipif(not live_db_enabled(), reason=LIVE_DB_REASON)

SOC = 1


@pytest.fixture()
def act(monkeypatch):
    psycopg2 = pytest.importorskip("psycopg2")
    from psycopg2.extras import RealDictCursor
    os.environ.setdefault("DATABASE_URL", "postgresql://{u}:{p}@{h}:5432/{d}?sslmode=disable".format(
        u=os.getenv("PGUSER"), p=os.getenv("PGPASSWORD"), h=os.getenv("PGHOST"), d=os.getenv("PGDATABASE")))
    from app.services import up_aoa_actions as mod

    conn = psycopg2.connect(host=os.getenv("PGHOST"), dbname=os.getenv("PGDATABASE"), user=os.getenv("PGUSER"),
                            password=os.getenv("PGPASSWORD"), sslmode=os.getenv("PGSSLMODE", "prefer"))
    cur = conn.cursor(cursor_factory=RealDictCursor)

    class _Tx:
        @staticmethod
        def _execute(sql, params=None, fetch_one=False, fetch_all=False):
            cur.execute(sql, params)
            if fetch_one:
                row = cur.fetchone()
                return dict(row) if row else None
            if fetch_all:
                return [dict(r) for r in cur.fetchall()]
            return None

    monkeypatch.setattr(mod, "db", _Tx)
    mod._cur = cur
    yield mod
    conn.rollback()
    conn.close()


def _apt(act):
    return act.db._execute("SELECT id FROM apartments WHERE society_id=%s AND active ORDER BY id LIMIT 1", (SOC,), fetch_one=True)["id"]


def _walk(component, ids):
    if component is None or isinstance(component, (str, int, float)):
        return
    if isinstance(component, (list, tuple)):
        for c in component:
            _walk(c, ids)
        return
    cid = getattr(component, "id", None)
    if isinstance(cid, str):
        ids.add(cid)
    _walk(getattr(component, "children", None), ids)


# ── renderer + wiring ─────────────────────────────────────────────────────────
def test_card_renders_with_real_data_and_every_callback_id_exists(act):
    import dash
    import dash_bootstrap_components as dbc
    from app.dash_apps.callbacks.up_compliance_callbacks import register_up_compliance_callbacks
    from app.dash_apps.pages.up_compliance_card import render_up_compliance_card

    card = render_up_compliance_card(act.load_card_data(SOC))
    ids = set()
    _walk(card, ids)

    app = dash.Dash(__name__, external_stylesheets=[dbc.themes.BOOTSTRAP], suppress_callback_exceptions=True)
    app.layout = card
    register_up_compliance_callbacks(app)
    wired = set()
    for spec in app.callback_map.values():
        wired.update(i["id"] for i in spec["inputs"] + spec.get("state", []) if isinstance(i["id"], str))
    for out in app.callback_map:
        wired.update(p.split(".")[0] for p in out.strip("..").split("...."))
    missing = {i for i in wired if i.startswith("upc-")} - ids
    assert not missing, f"callbacks reference ids the card does not render: {sorted(missing)}"
    assert len([k for k in app.callback_map if "upc-" in k]) >= 11


def test_card_shows_inactive_notice_when_regime_off(act):
    from app.dash_apps.pages.up_compliance_card import render_up_compliance_card, render_up_compliance_body
    data = act.load_card_data(999999)
    assert data["rules_on"] is False
    assert "not active" in str(render_up_compliance_card(data))
    assert render_up_compliance_body(data) is not None


# ── handlers ──────────────────────────────────────────────────────────────────
def test_filing_upsert_and_validation(act):
    ok, _ = act.save_filing(SOC, None, 2025, "2026-07-20", "2026-08-05", "2026-08-04", "ABC & Co", True, False)
    assert ok
    ok, _ = act.save_filing(SOC, None, 2025, "2026-07-21", "2026-08-06", None, "ABC & Co", True, True)      # update, not duplicate
    assert act.db._execute("SELECT count(*) AS n FROM aoa_statutory_filings WHERE society_id=%s AND fy_start_year=2025", (SOC,), fetch_one=True)["n"] == 1
    ok, msg = act.save_filing(SOC, None, 2025, "2026-08-10", "2026-08-01", None, None, False, False)
    assert not ok and "before" in msg
    assert act.save_filing(SOC, None, None, None, None, None, None, False, False)[0] is False


def test_billing_basis_guards_then_switches(act):
    act.db._execute("UPDATE apartments SET undivided_interest_pct=NULL WHERE society_id=%s", (SOC,))
    ok, msg = act.set_billing_basis(SOC, "undivided_interest", 90000)
    assert not ok and "100%" in msg                                           # refuses while percentages are missing
    assert act.fill_undivided_interest(SOC)[0]
    assert act.set_billing_basis(SOC, "undivided_interest", None)[0] is False   # budget required
    ok, _ = act.set_billing_basis(SOC, "undivided_interest", 90000)
    assert ok
    row = act.db._execute("SELECT billing_basis, common_expense_budget_monthly FROM apt_charges_fines_basis "
                          "WHERE society_id=%s AND apt_id IS NULL AND apt_status ORDER BY start_date DESC LIMIT 1", (SOC,), fetch_one=True)
    assert row["billing_basis"] == "undivided_interest" and float(row["common_expense_budget_monthly"]) == 90000.0
    assert act.set_billing_basis(SOC, "per_sqft", None)[0]


def test_transfer_then_nodues(act):
    apt = _apt(act)
    assert act.record_transfer(SOC, 1, 999999, "2026-09-10", 100, "a", "b")[0] is False        # not this society's flat
    ok, msg = act.record_transfer(SOC, 1, apt, "2026-09-10", 6_000_000, "Seller", "Buyer")
    assert ok and "30,000.00" in msg
    tid = act.db._execute("SELECT id FROM apartment_transfers WHERE society_id=%s ORDER BY id DESC LIMIT 1", (SOC,), fetch_one=True)["id"]
    assert act.set_nodues(SOC, tid, "nodues_requested_on", "2026-09-12")[0]
    status = act.load_card_data(SOC)["transfers"][0]
    assert status["nodues_status"] in ("pending", "deemed_granted")
    assert act.set_nodues(SOC, tid, "bogus", "2026-09-12")[0] is False
    assert act.set_nodues(SOC, 999999, "nodues_issued_on", "2026-09-12")[0] is False


def test_section22_cutoff_refused_until_every_step_done(act):
    apt = _apt(act)
    act.db._execute("UPDATE receivables SET status='paid', paid_amount=amount WHERE society_id=%s AND entity_id=%s AND role='apartment'", (SOC, apt))
    act.db._execute("""INSERT INTO receivables (society_id, entity_id, role, acc_id, description, base_amount, amount, due_date, status)
                       VALUES (%s,%s,'apartment',4210,'dues',9000,9000,'2026-01-01','pending')""", (SOC, apt))
    assert act.start_cutoff(SOC, None, apt, "water supply", "2026-01-01", None)[0]
    pid = act.db._execute("SELECT id FROM service_cutoff_proceedings WHERE society_id=%s ORDER BY id DESC LIMIT 1", (SOC,), fetch_one=True)["id"]
    ok, msg = act.set_cutoff_step(SOC, pid, "cut_off_on", "2026-09-20")
    assert not ok and "not met" in msg
    for step, day in (("notice_served_on", "2026-08-01"), ("gb_resolution_on", "2026-08-10"), ("copy_sent_to_authority_on", "2026-08-12"),
                      ("copy_sent_to_owner_on", "2026-08-12"), ("display_notice_on", "2026-08-13")):
        assert act.set_cutoff_step(SOC, pid, step, day)[0]
    ok, msg = act.set_cutoff_step(SOC, pid, "cut_off_on", "2026-09-05")
    assert not ok                                                              # one-month wait not over
    ok, _ = act.set_cutoff_step(SOC, pid, "cut_off_on", "2026-09-12")
    assert ok
    assert act.db._execute("SELECT status FROM service_cutoff_proceedings WHERE id=%s", (pid,), fetch_one=True)["status"] == "cut_off"
    assert act.start_cutoff(SOC, None, 999999, "water", "2026-01-01", None)[0] is False


def test_loan_and_repayment_round_trip(act):
    apt = _apt(act)
    ok, msg = act.disburse_loan(SOC, None, apt, "2026-09-01", 40000, 12, "bank", "roof", "")
    assert not ok and "resolution" in msg
    ok, msg = act.disburse_loan(SOC, None, apt, "2026-09-01", 40000, 12, "bank", "roof", "BM-9")
    assert ok
    loan = act.load_card_data(SOC)["loans"][0]
    assert float(loan["outstanding"]) == 40000.0 and loan["ledger_posted"]
    ok, _ = act.repay_loan(SOC, None, loan["id"], "2026-10-01", 15000, 400, "bank")
    assert ok
    assert float(act.load_card_data(SOC)["loans"][0]["outstanding"]) == 25000.0
    assert act.repay_loan(SOC, None, loan["id"], "2026-10-02", 99999, 0, "bank")[0] is False
    assert act.repay_loan(SOC, None, 999999, "2026-10-02", 1, 0, "bank")[0] is False


def test_bye_law7_and_annexure_export(act):
    rows, problem = act.bye_law7(SOC, "2026-05-10", "calendar_year")
    assert problem is None and len(rows) == act.db._execute("SELECT count(*) AS n FROM apartments WHERE society_id=%s AND active", (SOC,), fetch_one=True)["n"]
    assert act.bye_law7(SOC, None, None)[1]
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(act.annexure_workbook_bytes(SOC)))
    assert wb.sheetnames == ["Owners", "Loanees"] and wb["Owners"].max_row > 1


# ── session / role handling in the callback wrapper ───────────────────────────
def test_callbacks_refuse_non_admin_and_missing_society(act, monkeypatch):
    from app.dash_apps.callbacks import up_compliance_callbacks as cb
    monkeypatch.setattr(cb, "get_current_user_role", lambda: "apartment")
    toast, body = cb._run(lambda sid, uid: (True, "should not run"))
    assert "Admin only" in str(toast)
    monkeypatch.setattr(cb, "get_current_user_role", lambda: "admin")
    monkeypatch.setattr(cb, "get_current_society_id", lambda: None)
    assert "Society not resolved" in str(cb._run(lambda sid, uid: (True, "x"))[0])
    monkeypatch.setattr(cb, "get_current_society_id", lambda: SOC)
    monkeypatch.setattr(cb, "get_current_user_id", lambda: None)
    toast, body = cb._run(lambda sid, uid: (True, "saved"))
    assert "saved" in str(toast) and body is not None
    toast, body = cb._run(lambda sid, uid: (_ for _ in ()).throw(RuntimeError("boom")))
    assert "Could not save" in str(toast)
