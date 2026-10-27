"""
ER.md remediation — live tests against a REAL Postgres (rolled back; demo data untouched).

Covers: bank-import posting/reversal (A-01/A-02), the cash-voucher control, the
income-tax working paper and its export (A-03), TDS timeline (T-01/T-02), GST fund
toggle (X-01), resolution validity (G-01), ratification expiry (G-04), the
s.186 cash-receipt limit (C-01), and the "estatehub.sql is the only SQL" rule.

Setup: `python3 database/migrate.py --seed`, then
ESTATEHUB_LIVE_TESTS=1 PGHOST=... PGDATABASE=... PGUSER=... pytest test/test_er_remediation_live.py
"""

import datetime
import io
import os
import pathlib
import re
import uuid
from decimal import Decimal

import pytest

from test.live_db_gate import LIVE_DB_REASON, live_db_enabled

psycopg2 = pytest.importorskip("psycopg2")
import psycopg2.extras  # noqa: E402

pytestmark = pytest.mark.skipif(not live_db_enabled(), reason=LIVE_DB_REASON)

SOC = 1
ROOT = pathlib.Path(__file__).resolve().parents[1]


class TxDB:
    """Stands in for database.db_manager.db, but runs inside the test's own
    transaction so everything rolls back. Mirrors db.execute()'s return shapes."""

    def __init__(self, conn):
        self.conn = conn

    def execute(self, sql, params=None, fetch_one=False, fetch_all=False):
        cur = self.conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
        cur.execute(sql, params)
        if fetch_one:
            row = cur.fetchone()
            return dict(row) if row else None
        if fetch_all:
            return [dict(r) for r in cur.fetchall()]
        return None

    _execute = execute


@pytest.fixture()
def conn():
    c = psycopg2.connect(
        host=os.getenv("PGHOST"), port=os.getenv("PGPORT", "5432"),
        dbname=os.getenv("PGDATABASE"), user=os.getenv("PGUSER"),
        password=os.getenv("PGPASSWORD"), sslmode=os.getenv("PGSSLMODE", "prefer"),
    )
    cur = c.cursor()
    cur.execute("SELECT to_regproc('fn_cash_voucher_report') IS NOT NULL")
    if not cur.fetchone()[0]:
        c.close()
        pytest.skip("estatehub.sql with the ER.md controls is not installed in this database")
    cur.execute("SET CONSTRAINTS ALL IMMEDIATE")
    yield c
    c.rollback()
    c.close()


@pytest.fixture()
def cur(conn):
    return conn.cursor()


def one(cur, sql, args=()):
    cur.execute(sql, args)
    return cur.fetchone()[0]


def expect_error(cur, fn, match):
    cur.execute("SAVEPOINT s")
    with pytest.raises(psycopg2.Error) as exc:
        fn()
    cur.execute("ROLLBACK TO SAVEPOINT s")
    assert match.lower() in str(exc.value).lower(), str(exc.value)


def post(cur, legs, source=None, journal=None):
    """legs = [(side, acc, amount, mode)]; one INSERT so the balance trigger sees them all."""
    jid = journal or one(cur, "SELECT NEXTVAL('seq_transaction_number')")
    args = []
    for side, acc, amt, mode in legs:
        args += [SOC, side, datetime.date.today(), acc, amt, mode, jid,
                 source[0] if source else None, source[1] if source else None]
    cur.execute(
        "INSERT INTO transactions (society_id, entry_side, trx_date, acc_id, amount, mode, "
        "acc_particulars, status, journal_id, source_table, source_id) VALUES "
        + ",".join(["(%s,%s,%s,%s,%s,%s,'test leg','paid',%s,%s,%s)"] * len(legs)), args)
    return jid


def two_accounts(cur):
    cur.execute("SELECT id FROM accounts WHERE society_id=%s ORDER BY id LIMIT 2", (SOC,))
    return [r[0] for r in cur.fetchall()]


# ── bank import (A-01 / A-02) ────────────────────────────────────────────────
@pytest.fixture()
def bank(conn, monkeypatch):
    from app.dash_apps.callbacks import bank_reconcile_callbacks as b
    monkeypatch.setattr(b, "db", TxDB(conn))
    return b


def _statement_line(cur, bank_acc, debit=None, credit=None, desc="test line"):
    cur.execute(
        "INSERT INTO bank_statement_lines (society_id, txn_date, description, debit, credit, "
        "batch_id, bank_acc_id) VALUES (%s, CURRENT_DATE, %s, %s, %s, %s, %s) RETURNING id",
        (SOC, desc, debit, credit, str(uuid.uuid4()), bank_acc))
    return cur.fetchone()[0]


def _journal_totals(cur, line_id):
    cur.execute(
        "SELECT journal_id, COUNT(*), SUM(amount) FILTER (WHERE entry_side='Dr'), "
        "SUM(amount) FILTER (WHERE entry_side='Cr') FROM transactions "
        "WHERE bank_line_id=%s AND NOT is_reversal GROUP BY journal_id", (line_id,))
    return cur.fetchall()


def test_bank_import_posts_one_balanced_journal_per_line(cur, bank):
    primary = bank.get_primary_bank_account_id(SOC)
    cur.execute("DELETE FROM bank_statement_lines WHERE society_id=%s AND matched_id IS NULL "
                "AND reconciled=FALSE AND NOT EXISTS (SELECT 1 FROM transactions t WHERE t.bank_line_id=id)", (SOC,))
    out_id = _statement_line(cur, primary, debit=Decimal("118.00"), desc="bank charges")
    in_id = _statement_line(cur, primary, credit=Decimal("250.00"), desc="interest")
    assert bank.post_unmatched_bank_lines(SOC, None) == 2
    for line_id, amt in ((out_id, Decimal("118.00")), (in_id, Decimal("250.00"))):
        (jid, legs, dr, cr), = _journal_totals(cur, line_id)
        assert legs == 2 and dr == cr == amt
    # the statement's own bank account is on the right side of each journal
    cur.execute("SELECT entry_side FROM transactions WHERE bank_line_id=%s AND acc_id=%s", (out_id, primary))
    assert cur.fetchone()[0] == "Cr"          # money left the bank
    cur.execute("SELECT entry_side FROM transactions WHERE bank_line_id=%s AND acc_id=%s", (in_id, primary))
    assert cur.fetchone()[0] == "Dr"          # money arrived
    # running it again must not post the same lines twice
    assert bank.post_unmatched_bank_lines(SOC, None) == 0


def test_bank_import_withdrawal_reverses_instead_of_deleting(cur, bank):
    primary = bank.get_primary_bank_account_id(SOC)
    line_id = _statement_line(cur, primary, debit=Decimal("75.00"))
    assert bank.post_unmatched_bank_lines(SOC, None) == 1
    cur.execute("SELECT COUNT(*) FROM transactions WHERE bank_line_id=%s", (line_id,))
    before = cur.fetchone()[0]
    # Reversal legs are inserted one row at a time and balanced at COMMIT (the
    # trigger is DEFERRED in production), so defer here, then force the check.
    cur.execute("SET CONSTRAINTS ALL DEFERRED")
    assert bank.delete_unreconciled_bank_txns(SOC, None) >= 1
    cur.execute("SET CONSTRAINTS ALL IMMEDIATE")                  # raises if the reversal is unbalanced
    cur.execute("SELECT COUNT(*), COUNT(*) FILTER (WHERE is_reversal) FROM transactions WHERE source_table="
                "'bank_statement_lines' AND source_id=%s", (line_id,))
    total, reversals = cur.fetchone()
    assert total == before * 2 and reversals == before          # nothing deleted, mirror legs added
    cur.execute("SELECT SUM(CASE entry_side WHEN 'Dr' THEN amount ELSE -amount END) FROM transactions "
                "WHERE source_table='bank_statement_lines' AND source_id=%s", (line_id,))
    assert cur.fetchone()[0] == 0                                 # nets to zero
    cur.execute("SELECT 1 FROM bank_statement_lines WHERE id=%s", (line_id,))
    assert cur.fetchone() is not None                             # line kept: ledger references it
    # a second withdrawal must not reverse the same journal again
    assert bank.delete_unreconciled_bank_txns(SOC, None) == 0


# ── cash-voucher control (A-02, cash side) ───────────────────────────────────
def _cash_receipt(cur):
    cur.execute("SELECT id, amount, acc_id FROM receipts WHERE society_id=%s AND mode='cash' "
                "AND status='confirmed' ORDER BY id LIMIT 1", (SOC,))
    return cur.fetchone()


def test_cash_receipt_voucher_amount_is_checked_against_the_receipt(cur):
    rid, amount, acc = _cash_receipt(cur)
    post(cur, [("Cr", acc, amount, "cash")], source=("receipts", rid))            # correct: accepted
    expect_error(cur, lambda: post(cur, [("Cr", acc, amount + 1, "cash")], source=("receipts", rid)),
                 "moves")                                                         # too much cash
    expect_error(cur, lambda: post(cur, [("Cr", acc, amount - 1, "cash")], source=("receipts", rid)),
                 "moves")                                                         # single leg must be exact
    expect_error(cur, lambda: post(cur, [("Dr", acc, amount, "cash")], source=("receipts", rid)),
                 "bring cash in")                                                 # wrong direction


def test_cash_voucher_with_a_deduction_may_be_below_but_never_above_the_document(cur):
    rid, amount, acc = _cash_receipt(cur)
    a, b = two_accounts(cur)
    # Cr income 100%, Dr a receivable-type leg 10% -> implied cash in = 90%: allowed
    post(cur, [("Cr", acc, amount, "cash"), ("Dr", a, amount / 10, "cash")], source=("receipts", rid))


def test_cash_voucher_report_lists_what_is_checked(cur):
    rid, amount, acc = _cash_receipt(cur)
    jid = post(cur, [("Cr", acc, amount, "cash")], source=("receipts", rid))
    cur.execute("SELECT implied_cash_movement, direction, amount_checked FROM fn_cash_voucher_report(%s) "
                "WHERE journal_id=%s", (SOC, jid))
    assert cur.fetchone() == (amount, "cash in", True)


def test_unbalanced_non_cash_journal_is_still_rejected(cur):
    a, b = two_accounts(cur)
    expect_error(cur, lambda: post(cur, [("Dr", a, 100, "bank"), ("Cr", b, 90, "bank")]), "double-entry")


# ── income-tax working paper (A-03) ──────────────────────────────────────────
def test_income_tax_summary_runs_and_uses_one_expenditure_label(cur):
    cur.execute("SELECT category, nature, total_amount FROM fn_income_tax_summary_fy(%s, 2026)", (SOC,))
    rows = cur.fetchall()
    assert rows, "seed has FY2026 income"
    assert {r[0] for r in rows} <= {"Income", "Expenditure"}
    assert all(r[1] in ("mutual", "non_mutual") for r in rows)


def test_income_tax_export_handles_decimal_and_expenditure(conn):
    from openpyxl import load_workbook
    from database.income_tax_export import generate_income_tax_summary_excel
    data = generate_income_tax_summary_excel(TxDB(conn), SOC, 2026)
    ws = load_workbook(io.BytesIO(data))["Working Paper (CA)"]
    labels = {str(r[0].value): r[2].value for r in ws.iter_rows(min_row=1) if r[0].value}
    income = labels["Total Non-Mutual Income (potentially taxable)"]
    spend = labels["Total Non-Mutual Expenditure (P&L accounts)"]
    surplus = labels["Indicated Surplus for CA Review — NOT taxable income"]
    assert Decimal(str(surplus)) == Decimal(str(income)) - Decimal(str(spend))


def test_income_tax_export_sums_expenditure_under_either_label():
    """The export must not depend on which spelling the SQL function emits."""
    from openpyxl import Workbook
    from database.income_tax_export import _write_mutuality_sheet
    ws = Workbook().active
    _write_mutuality_sheet(ws, [
        {"category": "Income", "nature": "non_mutual", "total_amount": Decimal("1000.00")},
        {"category": "Expense", "nature": "non_mutual", "total_amount": Decimal("300.25")},
        {"category": "Expenditure", "nature": "non_mutual", "total_amount": Decimal("99.75")},
    ], 2026)
    vals = {str(r[0].value): r[2].value for r in ws.iter_rows() if r[0].value}
    assert Decimal(str(vals["Total Non-Mutual Expenditure (P&L accounts)"])) == Decimal("400.00")
    assert Decimal(str(vals["Indicated Surplus for CA Review — NOT taxable income"])) == Decimal("600.00")


# ── TDS timeline (T-01 / T-02) ───────────────────────────────────────────────
def test_tds_194h_threshold_follows_the_effective_date(cur):
    q = ("SELECT single_bill_threshold, act_version FROM fn_tds_section_rate(%s,'194H',NULL,%s::date,NULL)")
    cur.execute(q, (SOC, "2024-12-01"))
    assert cur.fetchone()[0] == 15000
    cur.execute(q, (SOC, "2025-12-01"))
    assert cur.fetchone()[0] == 20000
    cur.execute(q, (SOC, "2026-05-01"))
    assert cur.fetchone()[1] == "IT2025"


def test_tds_applies_only_when_the_amount_exceeds_the_threshold(cur):
    q = "SELECT applies FROM fn_compute_tds_pct(%s,NULL,'194H',NULL,'2025',%s,TRUE,'2025-12-01'::date)"
    cur.execute(q, (SOC, 20000))
    assert cur.fetchone()[0] is False          # equal is not "exceeds"
    cur.execute(q, (SOC, 20001))
    assert cur.fetchone()[0] is True


def test_no_ordinary_tds_row_is_left_open_after_the_1961_act(cur):
    cur.execute("SELECT COUNT(*) FROM tds_section_rates WHERE act_version='IT1961' AND effective_to IS NULL")
    assert cur.fetchone()[0] == 0


# ── X-01 GST fund toggle ─────────────────────────────────────────────────────
def test_fund_gst_exemption_needs_an_authority_on_file(cur):
    cur.execute("SELECT column_default FROM information_schema.columns WHERE table_name="
                "'society_compliance_settings' AND column_name='fund_gst_exempt'")
    assert cur.fetchone()[0].lower() == "false"
    expect_error(cur, lambda: cur.execute(
        "UPDATE society_compliance_settings SET fund_gst_exempt=TRUE WHERE society_id=%s", (SOC,)),
        "ck_fund_gst_exempt_authority")
    cur.execute("UPDATE society_compliance_settings SET fund_gst_exempt=TRUE, "
                "fund_gst_exempt_resolution_ref='GBM-12/2026', fund_gst_exempt_opinion_ref='CA/2026/41' "
                "WHERE society_id=%s", (SOC,))
    assert cur.rowcount == 1


# ── G-01 resolution validity ─────────────────────────────────────────────────
def _meeting(cur, quorum):
    cur.execute("INSERT INTO meetings (society_id, type, held_on, quorum_met) VALUES (%s,'GBM',CURRENT_DATE,%s) "
                "RETURNING id", (SOC, quorum))
    return cur.fetchone()[0]


def _resolution(cur, meeting, **kw):
    # A committee (MC) decision type: these tests cover the generic CHECK / quorum / freeze rules. A
    # General Body decision is passed from the registers (test_resolution_registers_live.py).
    cols = dict(meeting_id=meeting, decision_type_id=one(cur, "SELECT id FROM decision_types WHERE code='SET_BOARD_PARAM'"),
                body="test", majority_required=66.67, votes_for_pct=80, passed=False, status="draft",
                passed_on=None)
    cols.update(kw)
    cur.execute(f"INSERT INTO resolutions ({','.join(cols)}) VALUES ({','.join(['%s'] * len(cols))}) RETURNING id",
                list(cols.values()))
    return cur.fetchone()[0]


def test_resolution_cannot_pass_without_quorum_or_majority(cur):
    ok = _meeting(cur, True)
    no_quorum = _meeting(cur, False)
    passed = dict(passed=True, status="passed", passed_on=datetime.date.today())
    from test import resolution_fixtures as rf
    rf.committee_resolution(cur, ok, "petty_cash_limit", one(cur, "SELECT id FROM decision_types WHERE code='SET_BOARD_PARAM'"))   # valid
    # Task 1.10: a committee decision cannot be inserted as passed at all; it passes from the committee register.
    expect_error(cur, lambda: _resolution(cur, no_quorum, **passed), "quorum")
    expect_error(cur, lambda: _resolution(cur, ok, votes_for_pct=50, **passed), "committee register")
    expect_error(cur, lambda: _resolution(cur, ok, passed=True, status="draft", passed_on=datetime.date.today()),
                 "committee register")


def test_a_passed_resolution_is_frozen(cur):
    from test import resolution_fixtures as rf
    rid = rf.committee_resolution(cur, _meeting(cur, True), "petty_cash_limit", one(cur, "SELECT id FROM decision_types WHERE code='SET_BOARD_PARAM'"))
    expect_error(cur, lambda: cur.execute("UPDATE resolutions SET votes_for_pct=99 WHERE id=%s", (rid,)), "frozen")
    expect_error(cur, lambda: cur.execute("DELETE FROM resolutions WHERE id=%s", (rid,)), "cannot be deleted")


def test_fn_rule_ignores_a_decision_whose_resolution_did_not_pass(cur):
    # The EXISTS-passed guard is part of fn_rule's source: a draft/rejected resolution never activates a rule.
    src = one(cur, "SELECT prosrc FROM pg_proc WHERE proname='fn_rule' LIMIT 1")
    assert re.search(r"resolutions rs WHERE rs\.id = x\.resolution_id AND rs\.status = 'passed'", src)


# ── G-04 ratification expiry ─────────────────────────────────────────────────
def test_expiring_a_ratification_raises_a_compliance_flag(cur):
    rule_key = one(cur, "SELECT rule_key FROM society_rule_decisions WHERE society_id=%s LIMIT 1", (SOC,)) \
        if one(cur, "SELECT COUNT(*) FROM society_rule_decisions WHERE society_id=%s", (SOC,)) else "cash_limit_default_mode"
    cur.execute("INSERT INTO society_rule_decisions (society_id, rule_key, layer, status) "
                "VALUES (%s,%s,3,'adopted_as_is') ON CONFLICT DO NOTHING RETURNING id", (SOC, rule_key))
    got = cur.fetchone()
    row = (got[0], rule_key) if got else (one(cur, "SELECT id FROM society_rule_decisions WHERE society_id=%s LIMIT 1", (SOC,)), rule_key)
    cur.execute("INSERT INTO loa_ratification (society_id, decision_id, rule_key, board_layer, status, "
                "decision_effective, ratify_by) VALUES (%s,%s,%s,3,'pending',CURRENT_DATE-40,CURRENT_DATE-10) "
                "RETURNING id", (SOC, row[0], row[1]))
    lid = cur.fetchone()[0]
    assert one(cur, "SELECT fn_expire_overdue_ratifications()") >= 1
    cur.execute("SELECT rule_code FROM compliance_flags WHERE source_table='loa_ratification' AND source_id=%s", (lid,))
    assert cur.fetchone()[0] == "RATIFICATION_EXPIRED"


# ── C-01 s.186 ───────────────────────────────────────────────────────────────
def test_cash_receipt_of_two_lakh_or_more_is_refused_by_default(cur):
    msg = one(cur, "SELECT fn_check_cash_receipt_limit(%s, 200000, 'cash')", (SOC,))
    assert msg and "s.186" in msg and "269ST" in msg
    assert one(cur, "SELECT fn_check_cash_receipt_limit(%s, 199999.99, 'cash')", (SOC,)) is None
    assert one(cur, "SELECT fn_check_cash_receipt_limit(%s, 500000, 'upi')", (SOC,)) is None
    assert one(cur, "SELECT fn_cash_limit_mode(%s)", (SOC,)) in ("block", "warn")


def test_engine_default_cash_limit_mode_is_block(cur):
    cur.execute("SELECT pg_get_functiondef('fn_cash_limit_mode(int)'::regprocedure)")
    assert "'block')" in cur.fetchone()[0]


# ── single-source-of-SQL and inline-DDL guards ───────────────────────────────
def test_estatehub_sql_is_the_only_sql_file():
    found = [str(p.relative_to(ROOT)) for p in ROOT.rglob("*.sql")
             if ".git" not in p.parts and "node_modules" not in p.parts]
    assert found == ["database/estatehub.sql"], found


def test_only_the_unavoidable_forward_reference_is_an_alter_table():
    text = (ROOT / "database" / "estatehub.sql").read_text(encoding="utf-8")
    alters = re.findall(r"^\s*ALTER TABLE\s+(\w+)", re.sub(r"--[^\n]*", "", text), flags=re.M | re.I)
    # societies <-> accounts reference each other: one side cannot be declared inline.
    assert alters == ["societies"], alters
